#!/usr/bin/env python3
"""Validate Subjective Feedback Analysis 4.0.0 output without mandatory dependencies."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any


SKILL_DIR = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = SKILL_DIR.parents[2]
RUNTIME_DIR = WORKSPACE_ROOT / "shared-workflow/data-analysis/runtime"
if str(RUNTIME_DIR) not in sys.path:
    sys.path.insert(0, str(RUNTIME_DIR))
import contract_runtime
VERSION = "4.0.0"
TOP_LEVEL = {
    "run",
    "scope",
    "source_summary",
    "privacy",
    "data_quality",
    "analysis_results",
    "findings",
    "evidence_items",
    "knowledge_candidates",
    "limitations",
    "handoff",
}
FINAL_TOP_LEVEL = TOP_LEVEL | {"publication", "delivery"}
EVIDENCE_CLAIMS = {"user_report", "perception_signal", "service_signal", "sentiment_signal", "reported_error"}
COMPARISON_STATES = {
    "increased",
    "decreased",
    "stable",
    "newly_observed",
    "not_observed_in_current_sample",
    "shifted",
    "not_comparable",
}
REQUIRED_FORBIDDEN_USES = {
    "prevalence_claim",
    "causal_claim",
    "final_root_cause",
    "objective_impact_scale",
    "intervention_effectiveness",
    "knowledge_publish",
}
BANNED_EFFECT_STATUSES = {"effective", "ineffective", "resolved", "regressed", "harmful", "caused_by"}
REQUIRED_HANDOFF_FORBIDDEN_USES = {
    "treat_findings_as_governed_evidence",
    "publish_knowledge",
    "claim_prevalence",
    "claim_root_cause",
    "claim_effectiveness",
    "recalculate_visualization",
    "archive_without_governance",
}
BANNED_CLAIM_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"真实发生率",
        r"真实影响(人数|用户数)",
        r"已验证根因",
        r"最终根因",
        r"方案(已经)?有效",
        r"问题(已经)?解决",
        r"verified root cause",
        r"confirmed root cause",
        r"intervention is effective",
    )
]
SENSITIVE_PATTERNS = {
    "phone": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "email": re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE),
    "cn_id": re.compile(r"(?<!\d)\d{17}[0-9Xx](?!\d)"),
    "long_account": re.compile(r"(?<!\d)\d{16,19}(?!\d)"),
    "access_token": re.compile(r"(?i)\b(access[_-]?token|api[_-]?key|authorization)\s*[:=]\s*[^\s,;|]+"),
    "url_secret": re.compile(r"(?i)https?://[^\s]+[?&](token|access_token|key|signature|sig)=[^&\s]+"),
}
TEXT_ALIASES = {
    "customer_voice": ("feedback", "content", "description", "message", "工单内容", "问题描述", "反馈内容", "客户反馈", "用户反馈", "报错信息"),
    "support_restatement": ("客服转述", "问题摘要", "服务摘要", "处理描述"),
    "internal_context": ("处理记录", "跟进记录", "原因分析", "解决方案", "内部备注"),
}


def _fingerprint(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _redact(value: Any) -> tuple[str, list[str]]:
    text = str(value or "")
    risks = []
    replacements = {
        "phone": "[PHONE]",
        "email": "[EMAIL]",
        "cn_id": "[IDENTITY_NUMBER]",
        "long_account": "[ACCOUNT]",
        "access_token": "[CREDENTIAL]",
        "url_secret": "[URL_WITH_SECRET]",
    }
    for risk, pattern in SENSITIVE_PATTERNS.items():
        if pattern.search(text):
            risks.append(risk)
            text = pattern.sub(replacements[risk], text)
    text = re.sub(
        r"(?i)(user_?id|employee_?id|order_?id|ticket_?id)\s*[:=|：]\s*[^\s,;|]+",
        r"\1=[ID]",
        text,
    )
    return re.sub(r"\s+", " ", text).strip(), sorted(set(risks))


def _load_rows(source: dict[str, Any]) -> tuple[list[dict[str, Any]], bytes]:
    if source.get("content_ref"):
        raw = str(source["content_ref"]).encode("utf-8")
        return [{"content": source["content_ref"]}], raw
    path = Path(str(source.get("path_ref") or ""))
    if not path.is_absolute():
        path = SKILL_DIR.parents[2] / path
    raw = path.read_bytes()
    fmt = str(source.get("format") or path.suffix.lstrip(".")).lower()
    if fmt in {"csv", "tsv"}:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.DictReader(handle, delimiter="\t" if fmt == "tsv" else ",")), raw
    if fmt in {"json", "jsonl"}:
        if fmt == "jsonl":
            return [json.loads(line) for line in raw.decode("utf-8-sig").splitlines() if line.strip()], raw
        data = json.loads(raw.decode("utf-8-sig"))
        return data if isinstance(data, list) else data.get("records", data.get("rows", [])), raw
    if fmt == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX input requires optional openpyxl>=3.1") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        selected = source.get("sheets")
        if isinstance(selected, str):
            selected = [selected]
        worksheets = [
            workbook[name] for name in selected
        ] if selected else workbook.worksheets
        for sheet in worksheets:
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if header is None:
                continue
            fields = [str(value or f"unnamed_{index + 1}") for index, value in enumerate(header)]
            rows.extend(dict(zip(fields, values)) for values in iterator)
        workbook.close()
        return rows, raw
    raise ValueError(f"unsupported subjective input format: {fmt}")


STRUCTURED_LABELS = {
    "场景": ("场景", "问题场景", "descpScene"),
    "问题": ("问题", "问题描述", "descpQuestion"),
    "诉求": ("诉求", "用户诉求", "descpAppeal"),
    "其他风险": ("其他风险", "otherRisks"),
}


def _structured_parts(row: dict[str, Any]) -> list[tuple[str, str]]:
    parts = []
    for label, aliases in STRUCTURED_LABELS.items():
        value = next((row.get(key) for key in aliases if row.get(key) not in (None, "")), None)
        if value is not None:
            parts.append((label, str(value).strip()))
    return parts


def _ordered_text(row: dict[str, Any], fields: list[str]) -> tuple[str, list[dict[str, str]]]:
    values = []
    for field in fields:
        value = row.get(field)
        if value not in (None, ""):
            values.append({"field": field, "value": str(value).strip()})
    structured = _structured_parts(row)
    for label, value in structured:
        if not any(item["value"] == value for item in values):
            values.append({"field": label, "value": value})
    text = "\n".join(f"【{item['field']}】{item['value']}" for item in values)
    return text, values


def _nps_class(score: float) -> str:
    if score >= 9:
        return "promoter"
    if score <= 6:
        return "detractor"
    return "passive"


def _nps_fact(
    responses: list[dict[str, Any]],
    *,
    dimensions: dict[str, str],
    invited_count: int | None,
) -> dict[str, Any]:
    counts = {"promoter": 0, "passive": 0, "detractor": 0}
    for item in responses:
        counts[_nps_class(item["score"])] += 1
    response_count = len(responses)
    nps = (
        (counts["promoter"] - counts["detractor"]) / response_count * 100
        if response_count
        else None
    )
    return {
        "measurement_scope": {
            "grain": "survey_response",
            "dimensions": dimensions,
        },
        "response_count": response_count,
        "promoter_count": counts["promoter"],
        "passive_count": counts["passive"],
        "detractor_count": counts["detractor"],
        "nps": nps,
        "response_rate": {
            "status": "available" if invited_count is not None else "not_evaluable",
            "numerator": response_count,
            "denominator": invited_count,
            "value": (
                response_count / invited_count
                if invited_count not in (None, 0)
                else None
            ),
        },
    }


def analyze_nps_sources(
    source_rows: list[tuple[dict[str, Any], list[dict[str, Any]]]],
) -> dict[str, Any]:
    responses: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    invalid_score_count = 0
    duplicate_response_count = 0
    invited_total = 0
    invitation_known = False
    for source, rows in source_rows:
        bindings = source.get("field_bindings", {})
        score_field = bindings.get("nps_score")
        if not score_field:
            continue
        source_id = str(source.get("source_id") or "")
        response_id_field = bindings.get("response_id")
        invited_flag_field = bindings.get("invited_flag")
        if source.get("invited_count") is not None:
            invited_total += int(source["invited_count"])
            invitation_known = True
        elif invited_flag_field:
            invited_total += sum(
                str(row.get(invited_flag_field) or "").strip().lower()
                in {"1", "1.0", "true", "yes", "y", "是"}
                for row in rows
            )
            invitation_known = True
        for index, row in enumerate(rows, start=1):
            raw_score = row.get(score_field)
            if raw_score in (None, ""):
                continue
            try:
                score = float(raw_score)
            except (TypeError, ValueError):
                invalid_score_count += 1
                continue
            if score < 0 or score > 10:
                invalid_score_count += 1
                continue
            response_id = str(
                row.get(response_id_field) if response_id_field else index
            )
            key = (source_id, response_id)
            if key in seen:
                duplicate_response_count += 1
                continue
            seen.add(key)
            responses.append({
                "score": score,
                "business_board": (
                    str(row.get(bindings.get("business_board")) or "").strip()
                    if bindings.get("business_board")
                    else ""
                ),
                "enterprise_state": (
                    str(row.get(bindings.get("enterprise_state")) or "").strip()
                    if bindings.get("enterprise_state")
                    else ""
                ),
            })
    if not responses:
        return {
            "status": "not_evaluable",
            "reason": "nps_score_input_unavailable",
            "overall": None,
            "slices": [],
            "reliability": {
                "response_count": 0,
                "invalid_score_count": invalid_score_count,
                "duplicate_response_count": duplicate_response_count,
            },
            "limitations": [
                "NPS score field was not bound or contained no valid 0-10 responses"
            ],
        }
    overall = _nps_fact(
        responses,
        dimensions={},
        invited_count=invited_total if invitation_known else None,
    )
    grouped: dict[tuple[tuple[str, str], ...], list[dict[str, Any]]] = {}
    for response in responses:
        available = {
            key: response[key]
            for key in ("business_board", "enterprise_state")
            if response.get(key)
        }
        dimension_sets = [{key: value} for key, value in available.items()]
        if len(available) > 1:
            dimension_sets.append(available)
        for dimensions in dimension_sets:
            key = tuple(sorted(dimensions.items()))
            grouped.setdefault(key, []).append(response)
    return {
        "status": (
            "completed_with_caveats"
            if invalid_score_count or not invitation_known
            else "completed"
        ),
        "overall": overall,
        "slices": [
            _nps_fact(items, dimensions=dict(key), invited_count=None)
            for key, items in sorted(grouped.items())
        ],
        "reliability": {
            "response_count": len(responses),
            "invalid_score_count": invalid_score_count,
            "duplicate_response_count": duplicate_response_count,
            "response_deduplication": "source_id_plus_response_id_or_source_row",
            "inference_performed": False,
        },
        "limitations": (
            []
            if invitation_known
            else [
                "invited population denominator unavailable; response rate is not evaluable"
            ]
        ),
    }


def prepare_input(request: dict[str, Any]) -> dict[str, Any]:
    if request.get("artifact_type") != "subjective_analysis_request" or request.get("artifact_version") != VERSION:
        raise ValueError("request must use subjective_analysis_request 4.0.0")
    manifests, prepared = [], []
    nps_source_rows: list[tuple[dict[str, Any], list[dict[str, Any]]]] = []
    seen: set[str] = set()
    for source in request.get("input_sources", []):
        try:
            rows, raw = _load_rows(source)
            status, failure = "loaded", None
        except Exception as exc:
            rows, raw = [], b""
            status, failure = "unreadable", f"{type(exc).__name__}: {exc}"
        nps_source_rows.append((source, rows))
        bindings = source.get("field_bindings", {})
        source_type = str(source.get("source_type") or "")
        content_role = str(source.get("content_role") or ("internal_context" if source_type == "service_log" else "customer_voice"))
        configured_text_fields = bindings.get(content_role, [])
        if isinstance(configured_text_fields, str):
            configured_text_fields = [configured_text_fields]
        if not configured_text_fields and source_type != "service_log":
            configured_text_fields = list(TEXT_ALIASES.get(content_role, TEXT_ALIASES["customer_voice"]))
        usable_before = len(prepared)
        for index, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                continue
            raw_text, text_parts = _ordered_text(row, list(configured_text_fields))
            redacted, risks = _redact(raw_text)
            key = _fingerprint(redacted.encode("utf-8"))
            if not redacted or key in seen:
                continue
            seen.add(key)
            prepared.append(
                {
                    "prepared_item_id": f"prepared_{key[:14]}",
                    "source_id": source.get("source_id"),
                    "source_type": source_type,
                    "content_role": content_role,
                    "eligible_for_feedback_analysis": content_role != "internal_context",
                    "source_row": index,
                    "redacted_text": redacted,
                    "privacy_risk_types": risks,
                    "manual_review_required": bool(risks),
                    "channel": (
                        row.get(bindings.get("channel"))
                        if bindings.get("channel")
                        else source.get("channel") or source_type
                    ),
                    "product_location": {
                        "product_module": row.get(bindings.get("product_module")) if bindings.get("product_module") else None,
                        "page": row.get(bindings.get("page")) if bindings.get("page") else None,
                        "function": row.get(bindings.get("function")) if bindings.get("function") else None,
                        "operation_node": row.get(bindings.get("operation_node")) if bindings.get("operation_node") else None,
                        "mapping_basis": "explicit_field",
                        "mapping_confidence": "high",
                    },
                    "text_parts": [
                        {"field": item["field"], "redacted_value": _redact(item["value"])[0]}
                        for item in text_parts
                    ],
                    "segment_fields": {
                        "enterprise_id": row.get(bindings.get("enterprise_id")) if bindings.get("enterprise_id") else None,
                        "enterprise_name": row.get(bindings.get("enterprise_name")) if bindings.get("enterprise_name") else None,
                        **{
                            str(logical): row.get(source_field)
                            for logical, source_field in bindings.get("segment_fields", {}).items()
                        },
                    },
                }
            )
        usable_count = len(prepared) - usable_before
        analyzable_count = sum(
            item["eligible_for_feedback_analysis"]
            for item in prepared[usable_before:]
        )
        manifests.append(
            {
                "source_id": source.get("source_id"),
                "source_type": source_type,
                "content_role": content_role,
                "format": source.get("format"),
                "fingerprint": _fingerprint(raw) if raw else None,
                "row_count": len(rows),
                "usable_item_count": usable_count,
                "read_status": status,
                "analysis_status": (
                    "usable"
                    if analyzable_count
                    else "context_only"
                    if usable_count
                    else "unavailable"
                ),
                "unavailable_reason": (
                    failure
                    if failure
                    else "service_log_requires_explicit_text_field_binding"
                    if source_type == "service_log" and not configured_text_fields
                    else "no_usable_text_in_bound_fields"
                ) if not usable_count else None,
                "failure_reason": failure,
            }
        )
    analysis_item_count = sum(item["eligible_for_feedback_analysis"] for item in prepared)
    context_item_count = len(prepared) - analysis_item_count
    return {
        "artifact_type": "subjective_prepared_input",
        "artifact_version": VERSION,
        "project_name": request.get("project_name"),
        "source_mode": request.get("source_mode"),
        "analysis_mode": request.get("analysis_mode"),
        "source_manifest": manifests,
        "prepared_items": prepared,
        "nps_summary": analyze_nps_sources(nps_source_rows),
        "deduplicated_item_count": len(prepared),
        "analysis_item_count": analysis_item_count,
        "context_item_count": context_item_count,
        "raw_text_persisted": False,
        "status": (
            "ready"
            if analysis_item_count and all(item["analysis_status"] != "unavailable" for item in manifests if item["read_status"] == "loaded")
            else "ready_with_caveats"
            if analysis_item_count
            else "not_evaluable"
        ),
    }


def atomic_write(package: dict[str, Any], output: Path) -> Path:
    candidate = json.loads(json.dumps(package))
    candidate["run"]["write_status"] = "written"
    validate_output(candidate)
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".tmp", dir=output.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(candidate, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, output)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise
    return output


def _require_keys(value: dict[str, Any], keys: set[str], label: str) -> None:
    missing = keys - set(value)
    if missing:
        raise ValueError(f"{label} missing required fields: {', '.join(sorted(missing))}")


def _unique_ids(values: list[dict[str, Any]], key: str, label: str) -> set[str]:
    ids = [str(item.get(key) or "") for item in values]
    if any(not value for value in ids):
        raise ValueError(f"{label} contains an empty {key}")
    if len(ids) != len(set(ids)):
        raise ValueError(f"{label} contains duplicate {key}")
    return set(ids)


def _iter_text_fields(value: Any, path: str = ""):
    if isinstance(value, dict):
        for key, item in value.items():
            next_path = f"{path}.{key}" if path else key
            yield from _iter_text_fields(item, next_path)
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from _iter_text_fields(item, f"{path}[{index}]")
    elif isinstance(value, str):
        yield path, value


def _validate_sensitive_text(package: dict[str, Any]) -> None:
    for path, text in _iter_text_fields(package):
        for risk, pattern in SENSITIVE_PATTERNS.items():
            if pattern.search(text):
                raise ValueError(f"unredacted {risk} pattern found at {path}")


def _validate_claim_language(package: dict[str, Any]) -> None:
    statements = [str(item.get("statement") or "") for item in package["findings"] + package["evidence_items"]]
    for statement in statements:
        for pattern in BANNED_CLAIM_PATTERNS:
            if pattern.search(statement):
                raise ValueError(f"prohibited subjective claim language: {statement}")


def _validate_comparison(package: dict[str, Any]) -> None:
    comparison = package["analysis_results"].get("comparison")
    mode = package["run"].get("analysis_mode")
    if mode == "comparative_analysis" and not comparison:
        raise ValueError("comparative_analysis requires analysis_results.comparison")
    if mode == "discovery" and comparison is not None:
        raise ValueError("discovery must not emit a comparison result")
    if not comparison:
        return
    if comparison.get("baseline_type") == "run_local_comparison" and comparison.get("reusable") is not False:
        raise ValueError("run-local comparison baseline must not be reusable")
    for collection in ("theme_changes", "channel_changes"):
        for item in comparison.get(collection, []):
            status = item.get("status")
            if status not in COMPARISON_STATES or status in BANNED_EFFECT_STATUSES:
                raise ValueError(f"unsupported comparison status: {status}")


def _validate_nps(package: dict[str, Any]) -> None:
    summary = package["analysis_results"].get("nps_summary")
    if not isinstance(summary, dict):
        raise ValueError("analysis_results.nps_summary is required")
    if summary.get("status") == "not_evaluable":
        return
    facts = [summary.get("overall"), *summary.get("slices", [])]
    for fact in facts:
        if not isinstance(fact, dict):
            raise ValueError("NPS fact must be an object")
        _require_keys(
            fact,
            {
                "measurement_scope",
                "response_count",
                "promoter_count",
                "passive_count",
                "detractor_count",
                "nps",
                "response_rate",
            },
            "NPS fact",
        )
        total = (
            int(fact["promoter_count"])
            + int(fact["passive_count"])
            + int(fact["detractor_count"])
        )
        if total != int(fact["response_count"]):
            raise ValueError("NPS response_count does not reconcile")
        expected = (
            (fact["promoter_count"] - fact["detractor_count"]) / total * 100
            if total
            else None
        )
        if expected is not None and abs(float(fact["nps"]) - expected) > 1e-9:
            raise ValueError("NPS value does not match promoter/detractor counts")
        if fact.get("measurement_scope", {}).get("grain") != "survey_response":
            raise ValueError("NPS measurement_scope grain must be survey_response")


def _validate_references(package: dict[str, Any]) -> None:
    results = package["analysis_results"]
    item_ids = _unique_ids(results["normalized_feedback_items"], "item_id", "normalized feedback")
    clusters = results["theme_clusters"]
    cluster_ids = _unique_ids(clusters, "cluster_id", "theme clusters")
    finding_ids = _unique_ids(package["findings"], "finding_id", "findings")
    evidence_ids = _unique_ids(package["evidence_items"], "evidence_id", "evidence")
    candidate_ids = _unique_ids(package["knowledge_candidates"], "knowledge_candidate_id", "knowledge candidates")
    requests = results.get("objective_validation_requests", [])
    _unique_ids(requests, "objective_request_id", "objective validation requests")
    rule_signals = results.get("business_rule_signals", [])
    rule_ids = _unique_ids(rule_signals, "basis_id", "business rule signals")

    for basis in rule_signals:
        if basis.get("source_detail") in {"knowledge_candidate", "evidence_inferred", "unknown"} and basis.get("confirmation_status") == "confirmed":
            raise ValueError(f"business rule {basis['basis_id']} cannot be confirmed from its source")

    for screening in results.get("ticket_screening", []):
        if screening.get("source_item_id") not in item_ids:
            raise ValueError("ticket screening references an unknown item")
    for attention in results.get("analysis_attention_items", []):
        if attention.get("cluster_ref") not in cluster_ids:
            raise ValueError("analysis attention item references an unknown cluster")
    for sample_slice in results.get("sample_slices", []):
        if not set(sample_slice.get("theme_refs", [])) <= cluster_ids:
            raise ValueError("sample slice references an unknown theme cluster")
        if sample_slice.get("interpretation_boundary") != "current_feedback_sample_only_not_objective_prevalence":
            raise ValueError("sample slice must preserve the sample-only interpretation boundary")

    quote_ids: set[str] = set()
    for cluster in clusters:
        if not set(cluster.get("source_item_ids", [])) <= item_ids:
            raise ValueError(f"cluster {cluster['cluster_id']} references an unknown item")
        for quote in cluster.get("representative_quotes", []):
            quote_id = quote.get("quote_id")
            if not quote_id or quote_id in quote_ids:
                raise ValueError("quote ids must be non-empty and unique")
            quote_ids.add(quote_id)
            if quote.get("source_item_id") not in item_ids:
                raise ValueError(f"quote {quote_id} references an unknown item")

    eligible_findings = {item["finding_id"] for item in package["findings"] if item.get("evidence_eligible") is True}
    for finding in package["findings"]:
        if finding.get("claim_type") == "hypothesis" and finding.get("evidence_eligible") is not False:
            raise ValueError(f"hypothesis finding {finding['finding_id']} cannot be evidence-eligible")
        if finding.get("cluster_ref") not in cluster_ids or not set(finding.get("source_item_ids", [])) <= item_ids:
            raise ValueError(f"finding {finding['finding_id']} has an invalid cluster or item reference")

    evidence_count_by_finding: dict[str, int] = {}
    for evidence in package["evidence_items"]:
        if evidence.get("finding_ref") not in eligible_findings:
            raise ValueError(f"evidence {evidence['evidence_id']} must reference an evidence-eligible finding")
        finding_ref = str(evidence["finding_ref"])
        evidence_count_by_finding[finding_ref] = (
            evidence_count_by_finding.get(finding_ref, 0) + 1
        )
        if evidence.get("cluster_ref") not in cluster_ids:
            raise ValueError(f"evidence {evidence['evidence_id']} references an unknown cluster")
        if not set(evidence.get("source_item_ids", [])) <= item_ids:
            raise ValueError(f"evidence {evidence['evidence_id']} references an unknown item")
        if evidence.get("claim_type") not in EVIDENCE_CLAIMS:
            raise ValueError(f"unsupported evidence claim type: {evidence.get('claim_type')}")
        if evidence.get("privacy_status") not in {"safe", "restricted", "needs_review", "blocked"}:
            raise ValueError(f"evidence {evidence['evidence_id']} has invalid privacy_status")
        if not isinstance(evidence.get("manual_review_required"), bool):
            raise ValueError(f"evidence {evidence['evidence_id']} must declare manual_review_required")
        if evidence.get("privacy_status") in {"restricted", "needs_review", "blocked"} and evidence.get("manual_review_required") is not True:
            raise ValueError(f"evidence {evidence['evidence_id']} privacy restriction requires review")
        if not set(evidence.get("representative_quote_refs", [])) <= quote_ids:
            raise ValueError(f"evidence {evidence['evidence_id']} references an unknown quote")
        if not REQUIRED_FORBIDDEN_USES <= set(evidence.get("forbidden_uses", [])):
            raise ValueError(f"evidence {evidence['evidence_id']} is missing required forbidden uses")
        who = evidence.get("scene_descriptor", {}).get("who", {})
        if set(who) != {"actor_role", "user_segment", "stable_redacted_entity_ref"}:
            raise ValueError(f"evidence {evidence['evidence_id']} has an unsafe scene_descriptor.who")
        if evidence.get("lineage", {}).get("run_id") != package["run"].get("run_id"):
            raise ValueError(f"evidence {evidence['evidence_id']} lineage does not match the package run")
        basis = evidence.get("business_rule_basis")
        if basis is not None and basis.get("basis_id") not in rule_ids:
            raise ValueError(f"evidence {evidence['evidence_id']} references an undeclared business rule basis")

    for finding_id in sorted(eligible_findings):
        count = evidence_count_by_finding.get(finding_id, 0)
        if count != 1:
            raise ValueError(
                f"evidence-eligible finding {finding_id} must map to exactly one "
                f"evidence item; actual={count}"
            )

    for request in requests:
        if not set(request.get("related_finding_refs", [])) <= finding_ids:
            raise ValueError(f"objective request {request['objective_request_id']} references an unknown finding")
        if not set(request.get("related_evidence_refs", [])) <= evidence_ids:
            raise ValueError(f"objective request {request['objective_request_id']} references an unknown evidence item")
        if request.get("raw_feedback_included") is not False:
            raise ValueError(
                f"objective request {request['objective_request_id']} must not include raw feedback"
            )
        if not request.get("metric_refs") or not request.get("denominator_definition"):
            raise ValueError(
                f"objective request {request['objective_request_id']} must define metrics and denominator"
            )
        if not request.get("falsification_condition"):
            raise ValueError(
                f"objective request {request['objective_request_id']} must be falsifiable"
            )

    for candidate in package["knowledge_candidates"]:
        if candidate.get("manual_review_required") is not True:
            raise ValueError(f"knowledge candidate {candidate['knowledge_candidate_id']} must require review")
        if not set(candidate.get("source_item_ids", [])) <= item_ids:
            raise ValueError(f"knowledge candidate {candidate['knowledge_candidate_id']} references an unknown item")

    handoff = package["handoff"]
    expected_evidence_refs = [
        str(item["evidence_id"]) for item in package["evidence_items"]
    ]
    if handoff.get("evidence_refs", []) != expected_evidence_refs:
        raise ValueError(
            "handoff.evidence_refs must exactly match evidence_items in order"
        )
    if set(handoff.get("knowledge_review_request", {}).get("candidate_refs", [])) != candidate_ids:
        raise ValueError("knowledge review refs must exactly match knowledge candidates")


def _validate_handoff(package: dict[str, Any]) -> None:
    handoff = package["handoff"]
    if handoff.get("next_skill") != "evidence-govern":
        raise ValueError("default next skill must be evidence-govern")
    expected_payload_refs = [
        f"#/evidence_items/{index}"
        for index in range(len(package["evidence_items"]))
    ]
    if handoff.get("payload_refs", []) != expected_payload_refs:
        raise ValueError(
            "handoff.payload_refs must exactly match every evidence_items index"
        )
    knowledge = handoff.get("knowledge_review_request", {})
    if knowledge.get("auto_publish_allowed") is not False:
        raise ValueError("knowledge auto-publish must be disabled")
    if knowledge.get("candidate_refs") and knowledge.get("read_gate") != "governance_release_or_explicit_user_review":
        raise ValueError("knowledge candidates require governance release or explicit user review")
    visualization = handoff.get("visualization", {})
    if (
        visualization.get("target_skill") != "data-visualization"
        or "visualization" not in visualization.get("allowed_uses", [])
        or "recalculate_analysis" not in visualization.get("forbidden_uses", [])
        or not {"#/privacy", "#/limitations"}.issubset(
            set(visualization.get("payload_refs", []))
        )
    ):
        raise ValueError("visualization handoff is invalid")
    business = handoff.get("business_impact_access", {})
    if business.get("finding_maturity") != "problem_signal":
        raise ValueError("Business Impact findings must remain problem_signal")
    outcome = handoff.get("outcome_comparison_access", {})
    if outcome.get("requires_root_cause_and_intervention_contract") is not True:
        raise ValueError("Outcome comparison access requires root-cause and intervention contracts")
    if handoff.get("archive_access", {}).get("governance_gate_required") is not True:
        raise ValueError("archive access requires a governance gate")
    if not REQUIRED_HANDOFF_FORBIDDEN_USES <= set(handoff.get("forbidden_uses", [])):
        raise ValueError("handoff is missing required forbidden uses")


def _schema_validate(package: dict[str, Any]) -> bool:
    try:
        from jsonschema import Draft202012Validator, RefResolver
    except ImportError:
        return False
    schema = json.loads((SKILL_DIR / "schemas" / "output.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    resolver = RefResolver(
        base_uri=(SKILL_DIR / "schemas" / "output.schema.json").resolve().as_uri(),
        referrer=schema,
    )
    Draft202012Validator(schema, resolver=resolver).validate(package)
    return True


def validate_output(package: dict[str, Any]) -> bool:
    package_fields = frozenset(package)
    if package_fields not in {frozenset(TOP_LEVEL), frozenset(FINAL_TOP_LEVEL)}:
        raise ValueError(
            f"output top-level contract mismatch: expected {sorted(TOP_LEVEL)} "
            f"or {sorted(FINAL_TOP_LEVEL)}"
        )
    _require_keys(package["run"], {"run_id", "artifact_type", "artifact_version", "producer_skill", "source_mode", "analysis_mode", "status", "write_status"}, "run")
    run = package["run"]
    if run["artifact_type"] != "subjective_evidence_package" or run["artifact_version"] != VERSION:
        raise ValueError("output must use subjective_evidence_package 4.0.0")
    if run["producer_skill"] != "subjective-feedback-analysis":
        raise ValueError("unexpected producer_skill")
    if run["source_mode"] == "independent" and run.get("upstream_refs"):
        raise ValueError("independent mode cannot consume upstream refs")
    if package["privacy"].get("unredacted_text_persisted") is not False:
        raise ValueError("unredacted text must never be persisted")
    _validate_references(package)
    _validate_sensitive_text(package)
    _validate_claim_language(package)
    _validate_comparison(package)
    _validate_nps(package)
    _validate_handoff(package)
    return _schema_validate(package) if package_fields == frozenset(FINAL_TOP_LEVEL) else False


def render_markdown(package: dict[str, Any]) -> str:
    run = package["run"]
    purpose = run.get("analysis_purpose", {})
    results = package["analysis_results"]
    source_summary = package["source_summary"]
    lines = [
        "# 主观反馈分析报告",
        "",
        f"- 分析目的：{purpose.get('purpose_label', '本次反馈分析')}",
        f"- 项目：{run.get('project_name', '未声明')}",
        f"- 分析方式：{'前后反馈比较' if run.get('analysis_mode') == 'comparative_analysis' else '当前反馈发现'}",
        f"- 样本范围：{source_summary.get('valid_count', 0)} 条有效反馈（输入 {source_summary.get('input_count', 0)} 条）",
        "",
        "## 执行摘要",
        "",
    ]
    findings = package.get("findings", [])
    if findings:
        for item in findings:
            lines.append(
                f"- {item.get('statement', '未命名发现')}（样本内 {item.get('sample_count', 0)} 条）"
            )
    else:
        lines.append("- 当前样本未形成可发布的反馈发现。")
    lines.extend(["", "## 反馈主题与样本边界", ""])
    for cluster in results.get("theme_clusters", []):
        lines.append(
            f"### {cluster.get('label', '未命名主题')}\n\n"
            f"- 样本数：{cluster.get('sample_count', 0)}\n"
            f"- 样本占比：{cluster.get('sample_share', 0)}\n"
            f"- 置信度：{cluster.get('confidence', '未声明')}"
        )
        if cluster.get("bias_notes"):
            lines.append("- 偏差提醒：" + "；".join(cluster["bias_notes"]))
        lines.append("")
        if cluster.get("representative_quotes"):
            lines.append("- 代表性反馈：")
            lines.extend(
                f"  - “{quote.get('redacted_text', '')}”"
                for quote in cluster.get("representative_quotes", [])[:3]
            )
        if cluster.get("affected_pages") or cluster.get("affected_flows") or cluster.get("affected_journey_steps"):
            lines.append(
                "- 位置映射：页面 "
                + "、".join(cluster.get("affected_pages", []))
                + "；功能/流程 "
                + "、".join(cluster.get("affected_flows", []))
                + "；操作节点 "
                + "、".join(cluster.get("affected_journey_steps", []))
            )
        lines.append("")
    slices = results.get("sample_slices", [])
    if slices:
        lines.extend(["## 页面、功能与业务群体切片", ""])
        lines.append("- 以下占比只描述当前反馈样本，不代表客观问题发生率或总体人群影响。")
        for item in slices:
            dimensions = " × ".join(f"{key}={value}" for key, value in item.get("dimensions", {}).items())
            lines.append(
                f"- {dimensions}：{item.get('sample_count', 0)} 条，样本占比 {item.get('sample_share', 0):.1%}，"
                f"{'样本充足' if item.get('sample_status') == 'sufficient' else '低样本，仅供观察'}；"
                f"映射依据 {item.get('mapping_basis', 'unknown')}。"
            )
        lines.append("")
    nps_summary = results.get("nps_summary", {})
    if nps_summary.get("status") in {"completed", "completed_with_caveats"}:
        overall = nps_summary["overall"]
        lines.extend(["## NPS 与响应样本", ""])
        lines.append(
            f"- 当前有效响应 {overall['response_count']} 份；"
            f"推荐者 {overall['promoter_count']}、中立者 {overall['passive_count']}、"
            f"贬损者 {overall['detractor_count']}；NPS 为 {overall['nps']:.1f}。"
        )
        response_rate = overall["response_rate"]
        if response_rate.get("status") == "available":
            lines.append(
                f"- 响应率：{response_rate['numerator']}/{response_rate['denominator']}"
                f"（{response_rate['value']:.1%}）。"
            )
        else:
            lines.append("- 未提供受邀样本分母，本次不计算响应率。")
        for fact in nps_summary.get("slices", []):
            dimensions = " × ".join(
                f"{key}={value}"
                for key, value in fact["measurement_scope"]["dimensions"].items()
            )
            lines.append(
                f"- {dimensions}：{fact['response_count']} 份响应，NPS {fact['nps']:.1f}。"
            )
        lines.append("- NPS 只描述问卷响应样本；是否由体验优化导致，交由效果回检判断。")
        lines.append("")
    attention = results.get("analysis_attention_items", [])
    if attention:
        lines.extend(["## 分析关注项", ""])
        attention_labels = {
            "high_attention": "高度关注",
            "medium_attention": "中度关注",
            "low_attention": "一般观察",
            "needs_review": "需要人工复核",
        }
        for item in attention:
            lines.append(f"- {attention_labels.get(item.get('level'), item.get('level'))}：{'；'.join(item.get('basis', []))}")
        lines.append("")
    requests = results.get("objective_validation_requests", [])
    if requests:
        lines.extend(["## 需要客观数据验证的事项", ""])
        for item in requests:
            lines.append(
                f"- {item.get('question') or item.get('validation_question') or item.get('hypothesis') or '需要补充客观验证'}："
                f"指标 {', '.join(item.get('metric_refs', []))}；分母 {item.get('denominator_definition', '未声明')}。"
            )
        lines.append("")
    loss_findings = [
        item for item in package.get("findings", [])
        if item.get("business_loss_signal") == "explicitly_reported"
    ]
    if loss_findings:
        lines.extend(["## 用户明确表达的业务影响", ""])
        for item in loss_findings:
            lines.append(f"- {item.get('business_loss_description') or item.get('statement')}")
        lines.append("")
    comparison = results.get("comparison")
    if comparison:
        lines.extend(["## 前后反馈比较", ""])
        for collection in ("theme_changes", "channel_changes"):
            for item in comparison.get(collection, []):
                label = item.get("label") or item.get("theme") or item.get("channel") or "比较项"
                lines.append(f"- {label}：{item.get('status', 'not_comparable')}")
        if comparison.get("limitations"):
            lines.append("- 比较限制：" + "；".join(comparison["limitations"]))
        lines.append("")
    lines.extend(["## 数据质量与隐私边界", ""])
    quality = package["data_quality"]
    lines.extend(
        [
            f"- 数据状态：{quality.get('status', '未声明')}",
            f"- 去重后样本：{source_summary.get('valid_count', 0)}",
            "- 报告只说明当前样本中的反馈信号，不代表真实发生率、客观影响规模、根因或方案效果。",
        ]
    )
    if package["limitations"]:
        lines.extend(["", "## 限制", ""])
        lines.extend(f"- {item}" for item in package["limitations"])
    lines.extend(["", "## 下一步", ""])
    lines.append(
        "- 可交给证据治理进一步判断能否与其他来源共同使用；需要验证发生率或影响规模时，按客观验证请求补充数据。"
    )
    return "\n".join(lines).rstrip() + "\n"


def publish(
    package: dict[str, Any],
    request: dict[str, Any],
    workspace_root: Path,
    registry_path: Path,
) -> dict[str, Any]:
    prepared = contract_runtime.apply_request_envelope(package, request)
    prepared["run"]["write_status"] = "written"
    validate_output(prepared)
    with tempfile.TemporaryDirectory() as directory:
        temporary = Path(directory)
        candidate = temporary / "candidate.json"
        candidate.write_text(
            json.dumps(prepared, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        deliverables = []
        for item in request["requested_deliverables"]:
            bound = dict(item)
            if item["format"] == "markdown":
                report = temporary / "subjective-analysis-report.md"
                report.write_text(render_markdown(prepared), encoding="utf-8")
                bound["source_path"] = str(report)
            deliverables.append(bound)
        reference = contract_runtime.publish_result(
            workspace_root=workspace_root,
            registry_path=registry_path,
            artifact_id="subjective",
            candidate_path=candidate,
            purpose_key=request["analysis_purpose"]["purpose_key"],
            purpose_label=request["analysis_purpose"]["purpose_label"],
            deliverables=deliverables,
        )
    return reference


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command_or_package")
    parser.add_argument("package", nargs="?", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--request", type=Path)
    parser.add_argument("--workspace-root", type=Path, default=WORKSPACE_ROOT)
    parser.add_argument(
        "--registry",
        type=Path,
        default=WORKSPACE_ROOT / "shared-workflow/data-analysis/artifact-registry.json",
    )
    args = parser.parse_args()
    if args.command_or_package == "prepare-input":
        if args.package is None:
            parser.error("prepare-input requires a request path")
        request = json.loads(args.package.read_text(encoding="utf-8"))
        prepared = prepare_input(request)
        payload = json.dumps(prepared, ensure_ascii=False, indent=2) + "\n"
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(payload, encoding="utf-8")
        else:
            print(payload, end="")
        return 0
    path = args.package if args.command_or_package in {"validate", "write", "publish"} else Path(args.command_or_package)
    if path is None:
        parser.error("package path is required")
    package = json.loads(path.read_text(encoding="utf-8"))
    if args.command_or_package == "write":
        if args.output is None:
            parser.error("write requires an explicit --output candidate path")
        print(atomic_write(package, args.output))
        return 0
    if args.command_or_package == "publish":
        if args.request is None:
            parser.error("publish requires --request")
        request = json.loads(args.request.read_text(encoding="utf-8"))
        reference = publish(package, request, args.workspace_root, args.registry)
        print(json.dumps(reference, ensure_ascii=False, indent=2))
        return 0
    full_schema = validate_output(package)
    print(f"subjective-output=valid schema_validation={'full' if full_schema else 'core-only'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
