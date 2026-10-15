#!/usr/bin/env python3
"""Load, canonicalize, classify, and quality-check Objective event data."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import timezone
from pathlib import Path
from typing import Any, Iterable

import error_analyzer
from time_parsing import parse_time


SUPPORTED_FORMATS = {"csv", "tsv", "xlsx", "json", "jsonl"}
ALIASES = {
    "user_id": ["user_id", "userid", "userId", "用户id", "用户ID", "登录用户ID", "xft_login_id"],
    "enterprise_id": ["enterprise_id", "tenant_id", "corp_id", "企业id", "企业ID", "企业编号", "企业号", "xft_prj_code"],
    "enterprise_name": ["enterprise_name", "企业名称", "企业名", "xft_prj_name"],
    "task_instance_id": ["task_instance_id", "task_id", "taskId", "任务实例id", "任务ID", "用户任务埋点id", "xft_taskId"],
    "session_id": ["session_id", "sessionId", "会话id", "会话ID"],
    "event_time": ["event_time", "timestamp", "eventTime", "时间", "发生时间", "事件发生时间", "事件发生日期", "xft_time"],
    "event_name": ["event_name", "event", "eventName", "事件", "事件名称", "xft_eventID"],
    "event_custom_id": ["event_custom_id", "xft_cust_id"],
    "event_duration_ms": ["event_duration_ms", "xft_duration"],
    "journey_step": ["journey_step", "step", "node_name", "task_node_name", "步骤", "节点", "xft_taskNodeName"],
    "error_message": ["error_message", "message", "errorInfo", "错误信息", "报错信息", "xft_info"],
    "error_code": ["error_code", "code", "错误码", "xft_info3"],
    "error_presentation": ["error_presentation", "报错形式", "提示形式", "xft_info4"],
    "task_name": ["task_name", "任务名称", "任务名", "用户任务埋点名称"],
    "completion_flag": ["completion_flag", "is_completed", "completed", "是否完成", "完成标记"],
    "one_time_completion_flag": ["one_time_completion_flag", "first_time_completed", "是否一次完成", "是否一次性完成", "一次完成标记"],
    "task_duration": ["task_duration", "duration_seconds", "任务耗时", "任务时长", "总耗时(min)", "总耗时(分钟)"],
    "task_error_count": ["task_error_count", "error_count", "报错次数", "任务报错数", "报错总次数"],
    "journey_path": ["journey_path", "task_path", "任务路径", "旅程路径", "节点旅程"],
    "journey_step_order": ["journey_step_order", "step_order", "node_order", "节点顺序"],
    "journey_step_duration_seconds": ["journey_step_duration_seconds", "step_duration_seconds", "node_duration_seconds", "节点耗时(s)", "节点耗时(秒)"],
    "page": ["page", "page_name", "page_url", "页面", "页面名称", "页面地址", "页面路由地址"],
    "aggregate_completion_rate": ["aggregate_completion_rate", "completion_rate", "任务完成率"],
    "aggregate_one_time_completion_rate": ["aggregate_one_time_completion_rate", "one_time_completion_rate", "任务一次性完成率"],
    "aggregate_task_count": ["aggregate_task_count", "task_count", "任务总数", "总任务数"],
    "participant_count": ["participant_count", "总参与人数", "参与人数"],
    "aggregate_error_event_count": ["aggregate_error_event_count", "报错总次数"],
    "aggregate_error_affected_count": ["aggregate_error_affected_count", "报错影响总人数"],
    "error_is_complete": ["error_is_complete", "是否完整报错"],
    "business_scenario": ["business_scenario", "业务场景"],
    "sub_business_scenario": ["sub_business_scenario", "子业务场景名称"],
    "journey_name": ["journey_name", "旅程名称", "旅程名称taskname"],
    "registered_node_name": [
        "registered_node_name",
        "指定旅程关键节点名称",
        "指定旅程关键节点名称tasknodename",
    ],
    "registered_node_type": ["registered_node_type", "节点类型"],
    "registered_node_order": ["registered_node_order", "指定节点顺序", "旅程节点顺序"],
    "statistics_date": ["statistics_date", "统计日期", "数据日期"],
    "business_board": ["business_board", "business_module", "版块名称", "板块名称", "业务板块"],
    "business_board_level": ["business_board_level", "版块等级", "板块等级"],
    "enterprise_state": ["enterprise_state", "企业状态", "活跃状态"],
    "enterprise_state_member": ["enterprise_state_member", "是否属于该状态", "状态标记"],
}

SOURCE_SEMANTICS = {
    "aggregate_metric_summary",
    "task_summary",
    "journey_node_detail",
    "error_context_detail",
    "business_blueprint",
    "dimension_lookup",
    "enterprise_state_detail",
    "general_event_log",
}


def stable_id(prefix: str, *parts: Any) -> str:
    token = "|".join(str(part) for part in parts)
    return f"{prefix}_{hashlib.sha256(token.encode()).hexdigest()[:14]}"


WORKSPACE_ROOT = Path(__file__).resolve().parents[4]


def portable_path_ref(path: Path, source_id: str) -> str:
    """Return a workspace-relative path or the already-declared opaque source ID."""
    resolved = path.expanduser().resolve()
    try:
        return resolved.relative_to(WORKSPACE_ROOT).as_posix()
    except ValueError:
        return source_id


def time_sort_key(value: Any) -> tuple[int, float, str]:
    parsed = parse_time(value)
    raw = str(value or "")
    if parsed is None:
        return (1, 0.0, raw)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return (0, parsed.timestamp(), raw)


def _fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()[:16]


def _stringify(value: Any) -> Any:
    if value is None or (isinstance(value, float) and value != value):
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    return value


def _tag(
    rows: Iterable[dict[str, Any]],
    source_ref: str,
    sheet: str | None = None,
    start: int = 2,
    source_role: str = "event_data",
) -> list[dict[str, Any]]:
    output = []
    for row_number, row in enumerate(rows, start=start):
        item = {str(key): _stringify(value) for key, value in row.items()}
        item.update(
            {
                "_source_ref": source_ref,
                "_source_row": row_number,
                "_source_role": source_role,
            }
        )
        if sheet is not None:
            item["_source_sheet"] = sheet
        output.append(item)
    return output


def _load_delimited(
    path: Path, delimiter: str, source_ref: str, source_role: str
) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return _tag(
            csv.DictReader(handle, delimiter=delimiter),
            source_ref,
            source_role=source_role,
        )


def _load_json(
    path: Path, source_ref: str, jsonl: bool, source_role: str
) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig") as handle:
        if jsonl:
            records = [json.loads(line) for line in handle if line.strip()]
        else:
            data = json.load(handle)
            records = data if isinstance(data, list) else data.get("records", data.get("rows", []))
    if not isinstance(records, list) or any(not isinstance(item, dict) for item in records):
        raise ValueError(f"{path}: JSON input must be a list of objects or contain records/rows")
    return _tag(records, source_ref, source_role=source_role)


def _load_xlsx(
    path: Path, source_ref: str, sheets: list[str] | None, source_role: str
) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("XLSX input requires the optional dependency openpyxl>=3.1") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    selected = sheets or workbook.sheetnames
    output: list[dict[str, Any]] = []
    loaded: list[str] = []
    for sheet_name in selected:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{path}: worksheet not found: {sheet_name}")
        loaded.append(str(sheet_name))
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            continue
        fields = [str(value) if value is not None else f"unnamed_{index + 1}" for index, value in enumerate(header)]
        output.extend(
            _tag(
                (dict(zip(fields, values)) for values in iterator),
                source_ref,
                str(sheet_name),
                source_role=source_role,
            )
        )
    workbook.close()
    return output, loaded


def load_source(source: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    path = Path(source["path"]).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    fmt = str(source.get("format") or path.suffix.lstrip(".")).lower()
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError(f"unsupported input format: {fmt}; supported: {', '.join(sorted(SUPPORTED_FORMATS))}")
    source_ref = str(source.get("source_id") or path.name)
    source_role = str(source.get("role") or "event_data")
    sheets = source.get("sheets")
    if isinstance(sheets, str):
        sheets = [sheets]
    loaded_sheets: list[str] = []
    if fmt in {"csv", "tsv"}:
        rows = _load_delimited(
            path, "," if fmt == "csv" else "\t", source_ref, source_role
        )
    elif fmt in {"json", "jsonl"}:
        rows = _load_json(path, source_ref, fmt == "jsonl", source_role)
    else:
        rows, loaded_sheets = _load_xlsx(path, source_ref, sheets, source_role)
    fields = sorted({key for row in rows for key in row if not key.startswith("_source_")})
    manifest = {
        "source_id": source_ref,
        "role": source_role,
        "format": fmt,
        "path_ref": portable_path_ref(path, source_ref),
        "fingerprint": _fingerprint(path),
        "size_bytes": path.stat().st_size,
        "row_count": len(rows),
        "field_count": len(fields),
        "fields": fields,
        "sheets": loaded_sheets,
        "read_mode": "streaming_read_only" if fmt == "xlsx" else "full_json" if fmt == "json" else "iterative_read",
        "status": "loaded",
        "critical": bool(source.get("critical", False)),
    }
    return rows, manifest


def load_sources(sources: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    manifests: list[dict[str, Any]] = []
    for source in sources:
        try:
            loaded, manifest = load_source(source)
        except Exception as exc:
            if bool(source.get("critical", False)):
                raise
            manifest = {
                "source_id": str(source.get("source_id") or Path(str(source.get("path", "unknown"))).name),
                "role": source.get("role", "event_data"),
                "format": str(source.get("format") or Path(str(source.get("path", ""))).suffix.lstrip(".")).lower(),
                "path_ref": str(source.get("source_id") or Path(str(source.get("path", "unknown"))).name),
                "fingerprint": None,
                "size_bytes": None,
                "row_count": 0,
                "field_count": 0,
                "fields": [],
                "sheets": [],
                "read_mode": "unavailable",
                "status": "unreadable",
                "critical": False,
                "failure_reason": f"{type(exc).__name__}: {exc}",
            }
            loaded = []
        rows.extend(loaded)
        manifests.append(manifest)
    if not rows:
        failures = "; ".join(str(item.get("failure_reason") or item.get("source_id")) for item in manifests)
        raise RuntimeError(f"no readable input data sources: {failures}")
    return rows, manifests


def candidate_field_mapping(fields: Iterable[str]) -> dict[str, dict[str, Any]]:
    available = set(fields)
    lower = {field.lower(): field for field in available}
    result = {}
    for canonical, aliases in ALIASES.items():
        matches = list(dict.fromkeys([name for name in aliases if name in available] + [lower[name.lower()] for name in aliases if name.lower() in lower]))
        match = matches[0] if len(matches) == 1 else None
        result[canonical] = {
            "source_field": match,
            "candidate_fields": matches,
            "confidence": "high" if len(matches) == 1 else "ambiguous" if matches else "unmapped",
            "basis": "name_alias" if matches else "no_candidate",
        }
    return result


def canonicalize_rows(
    rows: list[dict[str, Any]],
    mappings: dict[str, dict[str, str | None]],
    fallbacks: dict[str, dict[str, list[str]]],
) -> list[dict[str, Any]]:
    output = []
    for row in rows:
        item = dict(row)
        source_ref = str(row.get("_source_ref") or "")
        mapping = mappings.get(source_ref, {})
        for canonical, source_field in mapping.items():
            item[canonical] = row.get(source_field, "") if source_field else ""
        for canonical, fields in fallbacks.get(source_ref, {}).items():
            if item.get(canonical) not in (None, ""):
                continue
            item[canonical] = next(
                (row.get(field) for field in fields if row.get(field) not in (None, "")),
                "",
            )
        if item.get("journey_step"):
            step, state, _ = error_analyzer.redact_message(item["journey_step"])
            item["journey_step"] = (
                step if state == "safe" else "unclassified_redacted_step"
            )
        output.append(item)
    return output


def infer_source_semantic_type(
    source: dict[str, Any],
    mapping: dict[str, str | None],
) -> tuple[str, str]:
    """Infer source capability from canonical field combinations, never filenames."""
    declared = str(source.get("semantic_type") or "")
    if declared:
        if declared not in SOURCE_SEMANTICS:
            raise ValueError(f"unsupported source semantic_type: {declared}")
        capability_checks = {
            "aggregate_metric_summary": bool(mapping.get("aggregate_completion_rate") or mapping.get("aggregate_one_time_completion_rate")),
            "task_summary": bool(mapping.get("task_instance_id") and (mapping.get("completion_flag") or mapping.get("one_time_completion_flag") or mapping.get("journey_path"))),
            "journey_node_detail": bool(mapping.get("task_instance_id") and mapping.get("journey_step")),
            "error_context_detail": bool(mapping.get("error_message") or mapping.get("error_code")),
            "general_event_log": bool(mapping.get("event_name") and (mapping.get("task_instance_id") or mapping.get("session_id") or mapping.get("user_id"))),
            "business_blueprint": bool(mapping.get("journey_name") and mapping.get("registered_node_name")),
            "enterprise_state_detail": bool(
                mapping.get("enterprise_id")
                and (
                    mapping.get("enterprise_state")
                    or mapping.get("business_board")
                    or source.get("enterprise_state_type")
                )
            ),
            "dimension_lookup": True,
        }
        if not capability_checks[declared]:
            raise ValueError(f"declared semantic_type {declared} is incompatible with available field capabilities")
        return declared, "explicit"
    role = str(source.get("role") or "")
    role_map = {
        "error_detail": "error_context_detail",
        "task_summary": "task_summary",
        "journey_node_detail": "journey_node_detail",
        "business_blueprint": "business_blueprint",
        "enterprise_state_detail": "enterprise_state_detail",
        "aggregate_metric_summary": "aggregate_metric_summary",
        "dimension_lookup": "dimension_lookup",
    }
    if role in role_map:
        return role_map[role], "role"
    if mapping.get("aggregate_completion_rate") or mapping.get("aggregate_one_time_completion_rate"):
        return "aggregate_metric_summary", "field_combination"
    if mapping.get("journey_name") and mapping.get("registered_node_name"):
        return "business_blueprint", "field_combination"
    if source.get("enterprise_state_type") and mapping.get("enterprise_id"):
        return "enterprise_state_detail", "field_combination"
    if mapping.get("enterprise_id") and any(
        "存续" in str(field) and str(field).startswith("是否")
        for field in source.get("_observed_fields", [])
    ):
        return "enterprise_state_detail", "field_combination"
    if mapping.get("completion_flag") or mapping.get("one_time_completion_flag") or mapping.get("journey_path"):
        return "task_summary", "field_combination"
    if mapping.get("journey_step") and mapping.get("journey_step_order") and mapping.get("task_instance_id"):
        return "journey_node_detail", "field_combination"
    if mapping.get("event_name") and mapping.get("event_time") and mapping.get("task_instance_id"):
        return "general_event_log", "field_combination"
    if (mapping.get("error_message") or mapping.get("error_code")) and mapping.get("task_instance_id"):
        return "error_context_detail", "field_combination"
    return "general_event_log", "field_combination"


def apply_dimension_joins(
    rows: list[dict[str, Any]],
    joins: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Apply explicit one-to-one lookup joins and fail on ambiguous cardinality."""
    results = []
    for join in joins:
        source_id = str(join["lookup_source_id"])
        target_key = str(join.get("target_key", "enterprise_id"))
        lookup_key = str(join.get("lookup_key", "enterprise_id"))
        value_fields = {
            str(target): str(source)
            for target, source in join.get("value_fields", {}).items()
        }
        lookup: dict[str, dict[str, Any]] = {}
        conflicts: set[str] = set()
        for row in rows:
            if str(row.get("_source_ref") or "") != source_id:
                continue
            key = str(row.get(lookup_key) or "")
            if not key:
                continue
            values = {target: row.get(source) for target, source in value_fields.items()}
            if key in lookup and lookup[key] != values:
                conflicts.add(key)
            lookup[key] = values
        if conflicts:
            raise ValueError(f"dimension lookup {source_id} has conflicting duplicate keys")
        eligible = matched = 0
        for row in rows:
            if row.get("_source_semantic_type") in {
                "dimension_lookup",
                "aggregate_metric_summary",
                "business_blueprint",
            }:
                continue
            key = str(row.get(target_key) or "")
            if not key:
                continue
            eligible += 1
            if key in lookup:
                matched += 1
                row.update(lookup[key])
        minimum = float(join.get("minimum_coverage", 0))
        coverage = matched / eligible if eligible else None
        results.append({
            "lookup_source_id": source_id,
            "target_key": target_key,
            "lookup_key": lookup_key,
            "lookup_key_count": len(lookup),
            "eligible_row_count": eligible,
            "matched_row_count": matched,
            "coverage": coverage,
            "status": (
                "completed"
                if lookup and (coverage is None or coverage >= minimum)
                else "not_evaluable"
            ),
        })
    return results


def _matches(row: dict[str, Any], conditions: list[dict[str, Any]]) -> bool:
    for condition in conditions:
        actual, operation = row.get(condition.get("field")), condition.get("op", "eq")
        if operation == "eq" and str(actual) != str(condition.get("value")):
            return False
        if operation == "in" and actual not in condition.get("values", []):
            return False
        if operation == "not_empty" and actual in (None, ""):
            return False
        if operation == "contains" and str(condition.get("value")) not in str(actual or ""):
            return False
    return True


def apply_exclusions(rows: list[dict[str, Any]], rules: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    kept = [row for row in rows if not any(_matches(row, rule.get("conditions", [])) for rule in rules)]
    return kept, len(rows) - len(kept)


def deduplicate_rows(rows: list[dict[str, Any]], contract: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], int]:
    contract = contract or {}
    event_fields = ("user_id", "task_instance_id", "session_id", "event_time", "event_name", "journey_step", "error_code", "error_message")
    blueprint_fields = (
        "business_scenario",
        "sub_business_scenario",
        "journey_name",
        "registered_node_name",
        "registered_node_type",
        "registered_node_order",
    )
    cross_source = bool(contract.get("cross_source_deduplication", {}).get("enabled", False))
    explicit_fields = contract.get("cross_source_deduplication", {}).get("event_key_fields", [])
    seen: set[tuple[str, ...]] = set()
    output = []
    for row in rows:
        is_blueprint = row.get("_source_semantic_type") == "business_blueprint"
        use_cross_source_event_key = (
            cross_source and bool(explicit_fields) and not is_blueprint
        )
        source_key = (
            ()
            if use_cross_source_event_key
            else (str(row.get("_source_ref") or ""),)
        )
        fields = (
            tuple(str(value) for value in explicit_fields)
            if use_cross_source_event_key
            else blueprint_fields
            if is_blueprint
            else event_fields
        )
        key = source_key + tuple(str(row.get(field) or "") for field in fields)
        if key not in seen:
            seen.add(key)
            output.append(row)
    return output, len(rows) - len(output)


def classify_events(
    rows: list[dict[str, Any]],
    contract: dict[str, Any],
    source_rules: dict[str, list[dict[str, Any]]] | None = None,
) -> dict[str, int]:
    source_rules = source_rules or {}
    event_mapping = dict(contract.get("event_mapping", {}))
    for category, field in (("success", "completion_events"), ("submit", "submit_events"), ("draft", "draft_events"), ("exit", "exit_events")):
        event_mapping.setdefault(category, contract.get(field, []))
    aliases = {"entry": "entry_event", "submit": "submit_event", "success": "success_event", "completion": "success_event", "exit": "exit_event", "draft": "draft_event", "backtrack": "backtrack_event", "help": "help_event", "error": "error_event"}
    inverse = {}
    for category, events in event_mapping.items():
        event_class = aliases.get(category, category if category.endswith("_event") else f"{category}_event")
        for event in events if isinstance(events, list) else [events]:
            inverse[str(event)] = event_class
    counts: Counter[str] = Counter()
    for row in rows:
        existing, name = str(row.get("event_class") or ""), str(row.get("event_name") or "")
        source_role_is_error = str(row.get("_source_role") or "") == "error_detail"
        rules = source_rules.get(str(row.get("_source_ref") or ""), [])
        matched_class = next(
            (
                str(rule.get("event_class"))
                for rule in rules
                if _matches(row, rule.get("conditions", []))
            ),
            "",
        )
        has_explicit_error_rule = any(
            str(rule.get("event_class")) == "error_event" for rule in rules
        )
        if existing:
            row["event_class"] = existing
        elif matched_class:
            row["event_class"] = matched_class
        elif source_role_is_error or (
            not has_explicit_error_rule
            and (row.get("error_message") or row.get("error_code"))
        ):
            row["event_class"] = "error_event"
        else:
            row["event_class"] = inverse.get(
                name,
                "step_event" if name or row.get("journey_step") else "noise_event",
            )
        counts[row["event_class"]] += 1
    return dict(counts)


def derive_task_instances(rows: list[dict[str, Any]], contract: dict[str, Any] | None = None) -> tuple[int, bool, list[str]]:
    contract = contract or {}
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if not row.get("task_instance_id"):
            group = str(row.get("session_id") or row.get("user_id") or "")
            if group:
                groups[group].append(row)
    derived = 0
    reliable = True
    reasons: list[str] = []
    restart_events = set(str(value) for value in contract.get("restart_events", []))
    timeout_seconds = contract.get("task_timeout_seconds")
    timeout_seconds = float(timeout_seconds) if timeout_seconds not in (None, "") else None
    for group, events in groups.items():
        events.sort(key=lambda row: time_sort_key(row.get("event_time")))
        attempt, active = 0, False
        prior_time: datetime | None = None
        for row in events:
            now = parse_time(row.get("event_time"))
            timed_out = bool(active and timeout_seconds is not None and prior_time and now and (now - prior_time).total_seconds() > timeout_seconds)
            explicit_restart = str(row.get("event_name") or "") in restart_events
            if not active or timed_out or explicit_restart:
                attempt += 1
                active = True
            row["task_instance_id"] = stable_id("derived_task", group, attempt)
            row["task_instance_derivation"] = "session_or_user_entry_terminal"
            derived += 1
            if row.get("event_class") in {"success_event", "exit_event", "draft_event"}:
                active = False
            if now:
                prior_time = now
        if events and not any(row.get("session_id") for row in events) and timeout_seconds is None and not restart_events:
            reliable = False
            reasons.append(f"user_only_task_boundary_without_timeout:{group}")
    if any(not row.get("task_instance_id") for row in rows):
        reliable = False
        reasons.append("rows_without_task_or_session_or_user_boundary")
    return derived, reliable, sorted(set(reasons))


def preprocess(sources: list[dict[str, Any]], contract: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    rows, manifests = load_sources(sources)
    source_by_id = {
        str(source.get("source_id") or Path(str(source["path"])).name): dict(source)
        for source in sources
    }
    global_explicit = contract.get("field_mapping", {})
    mappings: dict[str, dict[str, str | None]] = {}
    fallbacks: dict[str, dict[str, list[str]]] = {}
    source_candidates: dict[str, dict[str, dict[str, Any]]] = {}
    source_explicit: dict[str, dict[str, str]] = {}
    invalid_explicit: list[str] = []
    ambiguous_mappings: list[str] = []
    source_grains: dict[str, str] = {}
    source_semantics: dict[str, str] = {}
    source_semantic_basis: dict[str, str] = {}
    source_rules: dict[str, list[dict[str, Any]]] = {}
    for manifest in manifests:
        source_id = str(manifest["source_id"])
        source = source_by_id.get(source_id, {})
        fields = list(manifest.get("fields", []))
        source["_observed_fields"] = fields
        candidates = candidate_field_mapping(fields)
        explicit = {**global_explicit, **source.get("field_mapping", {})}
        source_explicit[source_id] = explicit
        mapping = {
            canonical: explicit.get(canonical) or candidate["source_field"]
            for canonical, candidate in candidates.items()
        }
        mapping.update(explicit)
        mappings[source_id] = mapping
        source_candidates[source_id] = candidates
        fallbacks[source_id] = {
            str(name): [str(value) for value in values]
            for name, values in source.get("field_fallbacks", {}).items()
        }
        if (
            mapping.get("journey_step") == "xft_taskNodeName"
            and mapping.get("event_name") == "xft_eventID"
        ):
            fallbacks[source_id].setdefault("journey_step", ["xft_eventID"])
        source_rules[source_id] = list(source.get("event_classification_rules", []))
        semantic_type, semantic_basis = infer_source_semantic_type(source, mapping)
        source_semantics[source_id] = semantic_type
        source_semantic_basis[source_id] = semantic_basis
        if not source_rules[source_id] and semantic_type == "error_context_detail":
            source_rules[source_id] = [
                {
                    "event_class": "error_event",
                    "conditions": [{"field": "error_message", "op": "not_empty"}],
                },
                {
                    "event_class": "error_event",
                    "conditions": [{"field": "error_code", "op": "not_empty"}],
                },
            ]
        if (
            not source_rules[source_id]
            and str(source.get("role") or "") == "error_detail"
        ):
            source_rules[source_id] = [
                {
                    "event_class": "error_event",
                    "conditions": [{"field": "error_message", "op": "not_empty"}],
                },
                {
                    "event_class": "error_event",
                    "conditions": [{"field": "error_code", "op": "not_empty"}],
                },
            ]
        invalid_explicit.extend(
            f"{source_id}:{canonical}"
            for canonical, source_name in explicit.items()
            if source_name not in fields
        )
        ambiguous_mappings.extend(
            f"{source_id}:{field}"
            for field, candidate in candidates.items()
            if candidate.get("confidence") == "ambiguous" and not explicit.get(field)
        )
        declared_grain = source.get("data_grain")
        if declared_grain:
            source_grains[source_id] = str(declared_grain)
        elif semantic_type in {"aggregate_metric_summary", "task_summary", "business_blueprint", "dimension_lookup", "enterprise_state_detail"}:
            source_grains[source_id] = "task_summary"
        else:
            has_summary_contract = any(
                mapping.get(name)
                for name in (
                    "completion_flag",
                    "one_time_completion_flag",
                    "task_error_count",
                    "journey_path",
                )
            )
            source_grains[source_id] = "task_summary" if has_summary_contract else "event"
        manifest["data_grain"] = source_grains[source_id]
        manifest["semantic_type"] = semantic_type
        manifest["semantic_type_basis"] = semantic_basis
        manifest["metric_authority"] = str(source.get("metric_authority") or "diagnostic")
        manifest["declared_scope"] = source.get("scope", {})
        if semantic_type == "enterprise_state_detail":
            manifest["enterprise_state_type"] = source.get("enterprise_state_type")
    canonical = canonicalize_rows(rows, mappings, fallbacks)
    for row in canonical:
        source_id = str(row.get("_source_ref") or "")
        row["_source_semantic_type"] = source_semantics.get(source_id, "general_event_log")
        source = source_by_id.get(source_id, {})
        if source.get("enterprise_state_type"):
            row["_enterprise_state_type"] = source["enterprise_state_type"]
        if source.get("blueprint_authority"):
            row["_blueprint_authority"] = source["blueprint_authority"]
    dimension_join_results = apply_dimension_joins(
        canonical, contract.get("dimension_joins", [])
    )
    rows_by_source: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in canonical:
        rows_by_source[str(row.get("_source_ref") or "")].append(row)
    for manifest in manifests:
        source_rows = rows_by_source.get(str(manifest["source_id"]), [])
        observed_times = sorted(
            parsed for parsed in (parse_time(row.get("event_time")) for row in source_rows)
            if parsed is not None
        )
        manifest["observed_time_window"] = {
            "start": observed_times[0].isoformat() if observed_times else None,
            "end": observed_times[-1].isoformat() if observed_times else None,
        }
        manifest["observed_grain_counts"] = {
            "task_instance": len({str(row.get("task_instance_id")) for row in source_rows if row.get("task_instance_id")}),
            "user": len({str(row.get("user_id")) for row in source_rows if row.get("user_id")}),
            "enterprise": len({str(row.get("enterprise_id")) for row in source_rows if row.get("enterprise_id")}),
        }
    input_count = len(canonical)
    canonical, excluded = apply_exclusions(canonical, contract.get("exclusion_rules", []))
    canonical, duplicates = deduplicate_rows(canonical, contract)
    analytic_rows = [
        row for row in canonical
        if row.get("_source_semantic_type") not in {
            "aggregate_metric_summary",
            "business_blueprint",
            "dimension_lookup",
            "enterprise_state_detail",
        }
    ]
    event_counts = classify_events(analytic_rows, contract, source_rules)
    derived, boundary_reliable, boundary_reasons = derive_task_instances(analytic_rows, contract)
    invalid_timestamps = sum(bool(row.get("event_time")) and parse_time(row.get("event_time")) is None for row in canonical)
    requested = set(contract.get("_requested_analyses", []))
    distinct_grains = set(source_grains.values())
    data_grain = next(iter(distinct_grains)) if len(distinct_grains) == 1 else "mixed"
    has_task_summary = "task_summary" in distinct_grains
    has_aggregate_metrics = "aggregate_metric_summary" in source_semantics.values()
    task_level_requested = requested & {"funnel", "journey", "errors", "segments"}
    required_fields = {"task_instance_id"} if task_level_requested or ("metrics" in requested and not has_aggregate_metrics) else set()
    if not has_task_summary and requested & {"metrics", "funnel", "feature_usage"}:
        required_fields.add("event_name")
    if not has_task_summary and requested & {"funnel", "journey"}:
        required_fields.add("journey_step")
    missing_core = [
        field for field in sorted(required_fields)
        if not any(row.get(field) not in (None, "") for row in canonical)
    ]
    source_failures = [item for item in manifests if item.get("status") == "unreadable"]
    error_sources = {
        source_id
        for source_id, mapping in mappings.items()
        if mapping.get("error_message") or mapping.get("error_code")
    }
    reliable_error_sources = {
        source_id
        for source_id in error_sources
        if source_semantics.get(source_id) == "error_context_detail"
        or any(
            str(rule.get("event_class")) == "error_event"
            for rule in source_rules.get(source_id, [])
        )
        or bool(contract.get("event_mapping", {}).get("error"))
    }
    error_classification_reliable = error_sources <= reliable_error_sources
    quality = {
        "status": "ready_with_caveats" if missing_core or invalid_timestamps or ambiguous_mappings or invalid_explicit or source_failures or not boundary_reliable or any(item["status"] == "not_evaluable" for item in dimension_join_results) else "ready",
        "input_row_count": input_count,
        "analysis_row_count": len(canonical),
        "excluded_row_count": excluded,
        "duplicate_row_count": duplicates,
        "derived_task_row_count": derived,
        "invalid_timestamp_count": invalid_timestamps,
        "missing_core_mappings": missing_core,
        "ambiguous_mappings": ambiguous_mappings,
        "invalid_explicit_mappings": invalid_explicit,
        "source_failures": source_failures,
        "task_boundary_reliable": boundary_reliable,
        "task_boundary_reasons": boundary_reasons,
        "event_class_counts": event_counts,
        "field_mapping": (
            {
                name: {
                    **source_candidates[next(iter(mappings))].get(name, {}),
                    "source_field": source,
                    **(
                        {"confidence": "explicit", "basis": "explicit_mapping"}
                        if name in source_explicit[next(iter(mappings))]
                        else {}
                    ),
                }
                for name, source in mappings[next(iter(mappings))].items()
            }
            if len(mappings) == 1
            else {}
        ),
        "source_field_mapping": {
            source_id: {
                name: {
                    **source_candidates[source_id].get(name, {}),
                    "source_field": source,
                    **(
                        {"confidence": "explicit", "basis": "explicit_mapping"}
                        if name in source_explicit[source_id]
                        else {}
                    ),
                }
                for name, source in mapping.items()
            }
            for source_id, mapping in mappings.items()
        },
        "source_data_grains": source_grains,
        "source_semantic_types": source_semantics,
        "source_semantic_type_basis": source_semantic_basis,
        "dimension_join_results": dimension_join_results,
        "data_grain": data_grain,
        "error_classification_reliable": error_classification_reliable,
        "row_level_persisted": False,
    }
    return canonical, quality, manifests
